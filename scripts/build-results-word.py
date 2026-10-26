"""Export the main analysis page with final figures and editable tables."""
from base64 import b64decode
from copy import deepcopy
import importlib.util
import os
from pathlib import Path
import re
import subprocess
import tempfile
from urllib.parse import urljoin

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from lxml import html

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / 'analysis/output/manuscript'
BASE = 'https://alisalman-et-al.github.io/arbiter-JAMIA/'


def normalized(text):
    return ' '.join(text.replace('\ufe0e', '').split())


def build():
    main = deepcopy(html.parse(str(ROOT / 'analysis/results.html')).xpath('//main')[0])
    for node in main.xpath('.//nav|.//script|.//style|.//a[contains(@class,"anchorjs-link")]'):
        node.getparent().remove(node)
    for paragraph in main.xpath('.//p[not(ancestor::table)]'):
        links = paragraph.xpath('./a/@href')
        if ('index.html' in links or any(link.endswith('Main-analysis.docx') for link in links)
                or any(re.search(r'/Figure[1-5][.]pdf$', link) for link in links)):
            paragraph.getparent().remove(paragraph)
    for anchor in main.xpath('.//a[@href]'):
        anchor.set('href', urljoin(BASE, anchor.get('href')))
    for heading in main.xpath('.//h3'):
        if re.fullmatch(r'Table [1-3]', normalized(heading.text_content())):
            heading.getparent().remove(heading)
    paragraphs = [normalized(p.text_content()) for p in main.xpath('.//p[not(ancestor::table)]|.//figcaption')]
    images = main.xpath('.//img')
    assert len(images) == 5
    with tempfile.TemporaryDirectory(prefix='arbiter-analysis-word-') as folder:
        folder = Path(folder)
        originals = []
        for i, image in enumerate(images, 1):
            body = b64decode(image.get('src').split(',', 1)[1])
            originals.append(body)
            path = folder / f'Figure{i}.png'
            path.write_bytes(body)
            image.set('src', str(path))
        source = folder / 'results.html'
        source.write_bytes(html.tostring(main, encoding='utf-8'))
        target = folder / 'Main-analysis.docx'
        subprocess.run(['quarto', 'pandoc', str(source), '--from=html', '--to=docx',
                        f'--output={target}'], check=True)
        doc = Document(target)
        assert len(doc.tables) == 3
        for i, table in enumerate(list(doc.tables), 1):
            native = Document(OUT / f'Table{i}.docx').tables[0]
            table._tbl.getparent().replace(table._tbl, deepcopy(native._tbl))
        for section in doc.sections:
            section.page_width, section.page_height = Inches(8.5), Inches(11)
            section.top_margin = section.bottom_margin = Inches(.65)
            section.left_margin = section.right_margin = Inches(.7)
            footer = section.footer.paragraphs[0]
            footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            field = OxmlElement('w:fldSimple')
            field.set(qn('w:instr'), 'PAGE')
            footer._p.append(field)
        for style in doc.styles:
            if style.type == 1:
                style.font.name = 'Arial'
                style.font.size = Pt(11)
                style.font.color.rgb = RGBColor(0, 0, 0)
                style.paragraph_format.line_spacing = 1.1
                style.paragraph_format.space_after = Pt(6)
        doc.styles['Title'].font.size = Pt(20)
        for level, size in [(1, 15), (2, 12)]:
            style = doc.styles[f'Heading {level}']
            style.font.size = Pt(size)
            style.font.bold = True
        for paragraph in doc.paragraphs:
            text = normalized(paragraph.text)
            if text == 'ARBITER main manuscript analysis':
                paragraph.style = doc.styles['Title']
            elif paragraph.style.name.startswith('Heading '):
                level = int(paragraph.style.name.split()[-1])
                paragraph.style = doc.styles[f'Heading {max(1, level - 1)}']
            if text in {'Figures', 'Tables'} or re.fullmatch(r'Figure [2-5]', text):
                paragraph.paragraph_format.page_break_before = True
            if re.match(r'Table [1-3][.]', text):
                paragraph.style = doc.styles['Heading 2']
                paragraph.paragraph_format.page_break_before = not text.startswith('Table 1.')
            if paragraph._p.xpath('.//w:drawing'):
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                paragraph.paragraph_format.keep_with_next = True
            if text.startswith('Figure ') and not re.fullmatch(r'Figure [1-5]', text):
                paragraph.paragraph_format.keep_together = True
        for shape, image in zip(doc.inline_shapes, images):
            scale = min(Inches(6.85) / shape.width, Inches(6.65) / shape.height)
            shape.width, shape.height = int(shape.width * scale), int(shape.height * scale)
            shape._inline.docPr.set('descr', image.get('alt', ''))
        doc.save(target)
        spec = importlib.util.spec_from_file_location('refresh', ROOT / 'scripts/report-refresh.py')
        refresh = importlib.util.module_from_spec(spec)
        spec.loader.exec_module(refresh)
        expected = [cells for i in range(1, 4) for cells in refresh.word_cells(OUT / f'Table{i}.docx')]
        assert refresh.word_cells(target) == expected, 'Exported table cells differ'
        written = Document(target)
        text = normalized(' '.join(p.text for p in written.paragraphs))
        for paragraph in paragraphs:
            assert paragraph in text, f'Exported paragraph missing: {paragraph[:80]}'
        assert len(written.inline_shapes) == 5
        for shape, original in zip(written.inline_shapes, originals):
            rid = shape._inline.xpath('.//a:blip')[0].get(qn('r:embed'))
            assert written.part.related_parts[rid].blob == original
        os.replace(target, OUT / 'Main-analysis.docx')
    print('PASS: main analysis Word export preserves page text, five figure images and three editable tables.')


if __name__ == '__main__':
    build()
