"""Focused manuscript generation tests; no analyses or publication writes."""
import importlib.util
from pathlib import Path
import unittest
import os
import ast
import re

spec = importlib.util.spec_from_file_location('builder', Path(__file__).with_name('build-manuscript.py'))
builder = importlib.util.module_from_spec(spec)
spec.loader.exec_module(builder)


class ManuscriptStructure(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        manuscript_dir = Path(__file__).resolve().parents[1] / 'docs/manuscript'
        spec = importlib.util.spec_from_file_location('numbers', manuscript_dir / 'verify-numbers.py')
        cls.numbers = importlib.util.module_from_spec(spec)
        spec.loader.exec_module(cls.numbers)
        # Load the actual formatting helpers without executing the publication build.
        from docx import Document
        from docx.oxml import OxmlElement
        from docx.oxml.ns import qn
        from docx.shared import Pt
        from docx.text.paragraph import Paragraph
        path = manuscript_dir / 'assemble-publication.py'
        tree = ast.parse(path.read_text(), filename=str(path))
        helpers = ast.Module(body=[node for node in tree.body if isinstance(node, ast.FunctionDef)
                                  and node.name in {'format_abstract', 'bind_reporting_reference'}], type_ignores=[])
        cls.formatting = dict(re=re, OxmlElement=OxmlElement, qn=qn, Pt=Pt, Paragraph=Paragraph)
        exec(compile(helpers, str(path), 'exec'), cls.formatting)
        cls.document = staticmethod(Document)

    def test_repeated_abstract_and_body_headings_are_scoped(self):
        text = ('# Title\n\n## Front-matter notes\n\n## Abstract\n'
                '### Objective\n\nAim.\n\n### Materials and Methods\n\nDesign.\n'
                '### Results\n\nFinding.\n\n### Discussion\n\nAbstract interpretation.\n'
                '### Conclusion\n\nAbstract conclusion.\n\n## Introduction\n'
                '### Background and significance\n\nContext.\n\n### Objective\n\nQuestion.\n'
                '## Materials and methods\n\nMethods.\n\n## Results\n\nResults.\n'
                '## Discussion\n\nBody interpretation.\n\n### Conclusion\n\nBody conclusion.\n'
                '## Back matter\n\nDeclarations.\n')
        self.assertEqual(self.numbers.structure_checks(text), [])
        counts, total = self.numbers.word_counts(text)
        self.assertEqual((counts['Abstract'], counts['Introduction'], counts['Discussion'], total), (7, 2, 4, 8))
        self.assertTrue(any('Introduction heading mismatch' in failure for failure in
                            self.numbers.structure_checks(text.replace('### Background and significance', '### Background'))))
        doc = self.document()
        for heading, style in [('ABSTRACT', 'Heading 1'), ('Discussion', 'Heading 2'),
                               ('Conclusion', 'Heading 2'), ('INTRODUCTION', 'Heading 1'),
                               ('Background and significance', 'Heading 2'), ('Objective', 'Heading 2'),
                               ('DISCUSSION', 'Heading 1'), ('Conclusion', 'Heading 2')]:
            doc.add_paragraph(heading, style)
            doc.add_paragraph('Prose.')
        self.formatting['format_abstract'](doc)
        self.assertEqual([p.text for p in doc.paragraphs if p.style.name == 'Heading 2'],
                         ['Background and significance', 'Objective', 'Conclusion'])
        self.assertEqual([p.text for p in doc.paragraphs if p.text.endswith('. Prose.')],
                         ['Discussion. Prose.', 'Conclusion. Prose.'])

    def test_reporting_reference_binding_survives_renumbering_and_run_split(self):
        doc = self.document()
        p = doc.add_paragraph('Reporting was guided by TRIPOD-LLM and GRRAS, with author-completion fields.')
        p.add_run('[24,25]')
        self.formatting['bind_reporting_reference'](p)
        self.assertTrue(p.text.endswith('fields.\u2060[24,25]'))
        other = doc.add_paragraph('Unrelated fields.[24,25]')
        self.formatting['bind_reporting_reference'](other)
        self.assertEqual(other.text, 'Unrelated fields.[24,25]')

    def test_command_metadata_does_not_exempt_adjacent_study_numbers(self):
        text = '`/rob2-assess` 987654321 /rob3-assess'
        spans = self.numbers.metadata_spans(text)
        classified = [m.group() for m in self.numbers.NUMBER.finditer(text)
                      if any(start <= m.start() < end for start, end, _ in spans)]
        self.assertEqual(classified, ['2'])

    def test_search_year_metadata_is_limited_to_eligibility_statement(self):
        text = 'Europe PMC for 2025–2026 meta-analyses; 2025 assessments and 2026 results.'
        spans = self.numbers.metadata_spans(text)
        classified = [m.group() for m in self.numbers.NUMBER.finditer(text)
                      if any(start <= m.start() < end for start, end, _ in spans)]
        self.assertEqual(classified, ['2025', '2026'])


class ManuscriptPage(unittest.TestCase):
    def setUp(self):
        self.text = '# Article\n\n## Abstract\n\nFull abstract.\n\nTable 1 and Table 2 appear together.\n\nLater Table 3.\n\n# References\n\n1. Complete reference.'
        self.tables = [f'<table><tr><td>cell {i}</td></tr></table>' for i in range(1, 4)]
        self.legends = [(f'Figure {i}. Legend <tag>.', 'Alt "quoted" & text') for i in range(1, 6)]

    def test_full_text_and_references_preserved_once(self):
        page = builder.manuscript_page(self.text, self.tables, self.legends)
        for paragraph in self.text.split('\n\n'):
            self.assertEqual(page.count(paragraph), 1)
        self.assertEqual(page.count('<table>'), 3)
        self.assertEqual(page.count('<figure '), 5)
        self.assertIn('alt="Alt &quot;quoted&quot; &amp; text"', page)
        self.assertIn('Legend &lt;tag&gt;.', page)
        for i in range(1, 6):
            self.assertIn(f'output/manuscript/Figure{i}.png', page)
            self.assertIn(f'output/manuscript/Figure{i}.pdf', page)

    def test_tables_follow_first_citation(self):
        page = builder.manuscript_page(self.text, self.tables, self.legends)
        self.assertLess(page.index('Table 1 and Table 2'), page.index('cell 1'))
        self.assertLess(page.index('cell 1'), page.index('cell 2'))
        self.assertLess(page.index('cell 2'), page.index('Later Table 3'))
        self.assertLess(page.index('cell 3'), page.index('# References'))

    def test_missing_table_citation_fails(self):
        with self.assertRaisesRegex(ValueError, 'Table 3 citation'):
            builder.manuscript_page(self.text.replace('Table 3', 'table three'), self.tables, self.legends)

    def test_legend_inventory_and_explicit_alt(self):
        text = '\n\n'.join(f'Figure {i}. Caption.\n\nAlt text: Actual alt {i}.' for i in range(1, 6))
        self.assertEqual(builder.metadata.main_legends(text)[3], ('Figure 4. Caption.', 'Actual alt 4.'))
        with self.assertRaisesRegex(ValueError, 'five ordered'):
            builder.metadata.main_legends(text.replace('Figure 3.', 'Figure 4.'))


@unittest.skipUnless(os.environ.get('ARBITER_CHECK_BUILT') == '1', 'Set ARBITER_CHECK_BUILT=1 after building the manuscript.')
class BuiltManuscript(unittest.TestCase):
    def test_final_cells_images_and_unique_anchors(self):
        from docx import Document
        from lxml import html
        page = html.parse(str(builder.ANALYSIS / 'manuscript.html'))
        normalize = lambda s: ' '.join(s.split())
        tables = page.xpath('//section[@class="manuscript-table"]//table')
        self.assertEqual(len(tables), 3)
        for i, table in enumerate(tables, 1):
            actual = [normalize(''.join(c.itertext())) for c in table.xpath('.//tr/*[self::td or self::th]')]
            expected = [normalize(' '.join(''.join(p.itertext('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}t')) for p in c.xpath('./w:p'))) for c in Document(builder.OUT / f'Table{i}.docx').tables[0]._tbl.xpath('./w:tr/w:tc')]
            self.assertEqual(actual, expected, f'Table {i} physical cells and notes')
        images = page.xpath('//figure[contains(@class,"manuscript-figure")]//img')
        self.assertEqual(len(images), 5)
        self.assertTrue(all(im.get('alt') and im.get('src').startswith('data:image/') for im in images))
        ids = page.xpath('//@id')
        self.assertEqual(len(ids), len(set(ids)), 'Duplicate contents destinations')


if __name__ == '__main__':
    unittest.main()
