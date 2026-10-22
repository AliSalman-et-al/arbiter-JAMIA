from pathlib import Path
from docx import Document
from lxml import html,etree
from zipfile import ZipFile
import csv,json,re,hashlib,subprocess
ROOT=Path(__file__).resolve().parents[4]; QA=Path(__file__).parent;OUT=ROOT/'analysis/output/manuscript';checks={}
norm=lambda s:' '.join(s.split())
def cells(t):return [norm(c.text) for r in t.rows for c in r.cells]
m=Document(OUT/'Manuscript.docx')
for i,t in enumerate(m.tables,1):assert cells(t)==cells(Document(OUT/f'Table{i}.docx').tables[0]),i
assert len(m.tables)==3
checks['embedded_main_tables']='all cells and notes exactly match standalone tables (whitespace normalized)'
for name in ['Manuscript','Table3']:
 t=' '.join(cells(Document(OUT/f'{name}.docx').tables[-1]))
 assert '14/14 (100%)' in t and '14/14 (100.0%)' not in t
 assert 'Unassisted rule-break cells show No./total (%; 95% CI); CI bounds are percentages.' in t
checks['Table3']='14/14 (100%) and explicit rule-break CI definition in both artifacts'
assert any(r.text=='P' and r.italic for p in m.paragraphs for r in p.runs if 'No P values' in p.text)
checks['P_style']='italic uppercase P values verified in actual Word runs'
supp=Document(OUT/'Supplement.docx');tabletexts=[' '.join(cells(t)) for t in supp.tables]
assert not any('paired differences are ARBITER minus the named comparator' in t for t in tabletexts)
assert any('subtract the second named source from the first' in t for t in tabletexts)
for t in tabletexts:
 if 'Unweighted Cohen κ against each independent expert' in t:
  assert 'Overall (all)' in t and 'Overall (observed)' in t and 'exact agreement (%)' not in t
 if 'Published is the reference. κ is unweighted;' in t:
  assert all(v not in t for v in ['AC2','sensitivity','specificity','All-Low','Pooled domains'])
checks['FVA01_03_VB01']='ordered contrasts; local 95% CI definition; tailored unweighted notes'
rows=list(csv.DictReader((ROOT/'analysis/output/tables/signalling-estimates.csv').open(encoding='utf-8-sig')))
zero=[r for r in rows if float(r['n'])==0];assert zero
assert all(float(r['count'])==0 for r in zero)
checks['zero_eligible_rows']=len(zero)
for label in ['eFigure 1.','eFigure 2.','eFigure 4.','eFigure 5.']:
 caps=[p.text for p in supp.paragraphs if p.text.startswith(label) and not p.style.name.lower().startswith(('heading','toc'))]
 assert caps and all(all(f'D{i},' in c for i in range(1,6)) for c in caps)
 if label!='eFigure 1.':assert all('Y, yes; PY, probably yes; N, no; PN, probably no; NI, no information.' in c and 'collapsed coding combines Y/PY and N/PN' in c for c in caps)
 if label=='eFigure 4.':assert all('eTable 13' in c and 'differ from the paired eligible intersections' in c for c in caps)
checks['VB03_04']='standalone domains/coding/collapse; precise eTable13 denominator locator'
# Compare the retained full trace, ignoring only added continuation labels.
paras=[p.text for p in supp.paragraphs if not p.text.startswith('Context: eh_')]
trace_text=norm(' '.join(paras));total=0
for f in (ROOT/'analysis/output/tables').glob('avatar-trace-*.csv'):
 for row in csv.DictReader(f.open(encoding='utf-8-sig')):
  for v in row.values():
   if v:assert norm(v) in trace_text,(f.name,v[:80]);total+=1
checks['verbatim_trace_cells']=total
metadata=[p for p in supp.paragraphs if p.text.startswith(('Source:','Page:','Lines:','Source-version hash:','Context:'))]
assert metadata and all(p.paragraph_format.keep_with_next for p in metadata)
checks['trace_metadata_context_keep_with_next']=len(metadata)
for name in ['results','supplement']:
 root=html.fromstring((ROOT/f'analysis/{name}.html').read_bytes())
 style='\n'.join(root.xpath('//style/text()'))
 assert 'padding-left: 8px !important' in style and 'text-align: right !important' in style,name
 checks[name+'_html_tables']=len(root.xpath('//table[contains(@class,"gt_table")]'))
before=json.loads((QA/'before-hashes.json').read_text());changed=[]
for f,h in before.items():
 p=ROOT/f
 if p.is_file() and hashlib.sha256(p.read_bytes()).hexdigest()!=h:changed.append(f)
checks['changed_output_files']=changed
raw=[f for f in changed if f.startswith('analysis/output/tables/') and '/display-' not in f]
assert not raw,raw
checks['raw_export_bytes_unchanged']=True
artwork=[f for f in before if re.search(r'/manuscript/(?:Figure[1245]|eFigure[^/]*[.](?:png|pdf|tiff)$)',f)]
assert all(hashlib.sha256((ROOT/f).read_bytes()).hexdigest()==before[f] for f in artwork)
checks['unchanged_artwork_files']=len(artwork)
# Baseline old PDF proves the original discrepancy.
text=lambda p:subprocess.check_output(['C:/Program Files/Calibre2/app/bin/pdftotext.exe','-layout',str(p),'-']).decode()
old=QA.parent/'remediation/proof-jama'
checks['baseline_stale_table_confirmed']=bool(re.search(r'14/14\s*\(100[.]0%\)',text(old/'Manuscript.pdf')) and re.search(r'14/14\s*\(100%\)',text(old/'Table3.pdf')))
checks['baseline_pdf_table3_rounding']={'Manuscript_100':bool(re.search(r'14/14\s*\(100%\)',text(old/'Manuscript.pdf'))),'Table3_100':bool(re.search(r'14/14\s*\(100%\)',text(old/'Table3.pdf')))}
(QA/'focused-checks.json').write_text(json.dumps(checks,indent=2));print(json.dumps(checks,indent=2))
