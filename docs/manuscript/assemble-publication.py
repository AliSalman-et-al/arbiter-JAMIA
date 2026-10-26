"""Finalize pagination and heading navigation for retained publication materials."""
from pathlib import Path
import re
from docx import Document
from docx.shared import Pt
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.style import WD_STYLE_TYPE

OUT = Path(__file__).resolve().parents[2] / 'analysis/output/manuscript'


def page_numbers(doc):
    for section in doc.sections:
        p = section.footer.paragraphs[0]; p.alignment = 2
        p.clear(); p.add_run('Page ')
        field = OxmlElement('w:fldSimple'); field.set(qn('w:instr'),'PAGE')
        p._p.append(field)
        for r in p.runs:r.font.name='Arial';r.font.size=Pt(9)

def bookkeep(doc):
    # Word updates TOC/PAGEREF fields on opening or with the COM proof helper.
    toc_fields=[f for f in doc.element.xpath('.//w:instrText') if f.text and f.text.strip().startswith('TOC ')]
    if toc_fields:
        if 'TOC 1' not in doc.styles:
            doc.styles.add_style('TOC 1',WD_STYLE_TYPE.PARAGRAPH,builtin=True)
        for style in doc.styles:
            if style.type==WD_STYLE_TYPE.PARAGRAPH and style.name.lower().startswith('toc'):
                style.font.name='Arial';style.font.size=Pt(10)
        # officer's renamed heading styles do not reliably populate Word's
        # built-in-heading TOC. Select actual stable reader locators explicitly.
        nav_name='Heading navigation'
        if nav_name not in doc.styles:
            nav=doc.styles.add_style(nav_name,WD_STYLE_TYPE.PARAGRAPH)
            nav.base_style=doc.styles['heading 2']
        for p in doc.paragraphs:
            if p.style.name.lower().startswith('heading'):
                selected=bool(re.match(r'^(eTable \d+[A-C]?[ .]|eFigure [1-5][ .]|[1-7]\. |eMethods|AVATAR complete trace$|Prompt templates and run provenance$)',p.text))
                if selected and '(continued ' not in p.text:p.style=doc.styles[nav_name]
        for field in toc_fields:field.text='TOC \\t "Heading navigation,1" \\h \\z'
    settings=doc.settings.element
    old=settings.find(qn('w:updateFields'))
    if old is None:old=OxmlElement('w:updateFields');settings.append(old)
    old.set(qn('w:val'),'true')
    ours=[b for b in doc.element.xpath('.//w:bookmarkStart') if b.get(qn('w:name'),'').startswith('loc_')]
    identifiers={b.get(qn('w:id')) for b in ours}
    for b in ours+[b for b in doc.element.xpath('.//w:bookmarkEnd') if b.get(qn('w:id')) in identifiers]:
        b.getparent().remove(b)
    next_id=max([int(b.get(qn('w:id'),'0')) for b in doc.element.xpath('.//w:bookmarkStart')]+[0])+1
    for i,p in enumerate(doc.paragraphs):
        if p.style.name.lower().startswith('heading'):
            start=OxmlElement('w:bookmarkStart');start.set(qn('w:id'),str(next_id+i));start.set(qn('w:name'),f'loc_{i}')
            end=OxmlElement('w:bookmarkEnd');end.set(qn('w:id'),str(next_id+i))
            p._p.insert(0,start);p._p.append(end)
    for p in doc.paragraphs:
        if p.style.name.lower().startswith('heading') or re.match(r'^e?Figure [1-5] ',p.text):
            for run in p.runs:
                run.text=re.sub(r'^(e?Table [0-9]+[A-C]?|e?Figure [0-9]+) ',r'\1. ',run.text)
    page_numbers(doc)


def finalize():
    for name in ['Table1','Table2','Table3','Supplement','TRIPOD-LLM']:
        path = OUT / f'{name}.docx'
        doc = Document(path)
        if name == 'Supplement':
            assert not doc.inline_shapes, 'The compact supplement has no supplementary figures'
            for p in doc.paragraphs:
                if not p.style.name.lower().startswith('heading'):
                    p.paragraph_format.line_spacing = 1.3
                    p.paragraph_format.space_after = Pt(4)
        bookkeep(doc)
        doc.save(path)
    print('Finalized three main tables, supplement and reporting map.')


if __name__ == '__main__':
    finalize()
