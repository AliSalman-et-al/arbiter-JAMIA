"""Check fixed study outputs and manuscript table/citation lineage."""
from pathlib import Path
import csv
import hashlib
import json
import re
from docx import Document
from lxml import html

ROOT = Path(__file__).resolve().parents[2]
HERE = Path(__file__).resolve().parent
baseline = json.loads((HERE / 'baseline-hashes.json').read_text())
changed = [name for name, digest in baseline.items()
           if hashlib.sha256((ROOT / name).read_bytes()).hexdigest() != digest]
assert not changed, f'Protected outputs changed: {changed}'
text = (ROOT / 'docs/manuscript/manuscript.md').read_text()
references = (ROOT / 'docs/manuscript/references.md').read_text()
order = []
for callout in re.findall(r'\[([\d,–-]+)\]', text):
    for group in callout.split(','):
        ends = re.split('[–-]', group)
        for number in range(int(ends[0]), int(ends[-1]) + 1):
            if number not in order:
                order.append(number)
listed = [int(n) for n in re.findall(r'^(\d+)\.', references, re.M)]
assert order == listed == list(range(1, len(listed) + 1)), (order, listed)
normalize = lambda s: ' '.join(s.split())
main = Document(ROOT / 'analysis/output/manuscript/Manuscript.docx')
page = html.fromstring((ROOT / 'analysis/manuscript.html').read_bytes())
web_tables = page.xpath('//section[contains(@class,"manuscript-table")]//table')
assert len(main.tables) == len(web_tables) == 3
cells_checked = 0
for i in range(1, 4):
    standalone = Document(ROOT / f'analysis/output/manuscript/Table{i}.docx').tables[0]
    source_cells = [normalize(c.text) for r in standalone.rows for c in r.cells]
    embedded_cells = [normalize(c.text) for r in main.tables[i-1].rows for c in r.cells]
    assert source_cells == embedded_cells, f'Table {i}: Word mismatch'
    web_text = normalize(web_tables[i-1].text_content())
    with (ROOT / f'analysis/output/tables/display-table{i}.csv').open(newline='', encoding='utf-8-sig') as stream:
        for row in csv.reader(stream):
            for value in row:
                if not value:
                    continue
                value = normalize(value)
                assert value in source_cells, (i, 'CSV value missing in Word', value)
                assert value in web_text, (i, 'CSV value missing in HTML', value)
                cells_checked += 1
assert len(page.xpath('//h1')) == 1, 'Manuscript must have one document title'
assert len(page.xpath('//figure[contains(@class,"manuscript-figure")]/img')) == 5
for img in page.xpath('//figure[contains(@class,"manuscript-figure")]/img'):
    assert len(img.get('alt', '')) > 50
print(json.dumps({'status': 'PASS', 'unchanged_sources': len(baseline),
                  'references_in_first_use_order': len(listed),
                  'display_csv_cells_in_both_formats': cells_checked,
                  'embedded_word_tables_equal_standalone': 3,
                  'figures_with_alt_text': 5}, indent=2))
